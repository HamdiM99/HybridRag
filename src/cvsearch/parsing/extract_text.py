"""Extraction de texte multi-format (PDF, DOCX, PPTX).

Principe : on essaie toujours l'extraction texte d'abord (rapide, gratuite).
Si le résultat est trop court par rapport au nombre d'images embarquées, on
considère que le contenu utile est dans une image (cas observé : CV où tout le
corps "WORK EXPERIENCE" est une image PNG collée dans un Word) et on bascule
en mode vision (voir vision.py) plutôt que d'indexer un profil quasi vide.
"""
from dataclasses import dataclass
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation

from cvsearch.config import get_settings


@dataclass
class ExtractionResult:
    text: str
    mode: str  # "text" ou "needs_vision"
    image_paths: list[Path]  # images à envoyer au modèle vision si mode == needs_vision


def _extract_pdf(path: Path) -> tuple[str, int]:
    doc = fitz.open(path)
    text_parts = []
    n_images = 0
    for page in doc:
        # "blocks" respecte mieux les mises en page multi-colonnes que le texte brut.
        blocks = page.get_text("blocks")
        blocks_sorted = sorted(blocks, key=lambda b: (round(b[1] / 20), b[0]))
        text_parts.append("\n".join(b[4].strip() for b in blocks_sorted if b[4].strip()))
        n_images += len(page.get_images(full=True))
    return "\n\n".join(text_parts), n_images


def _extract_docx(path: Path) -> tuple[str, int]:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))

    # Zones de texte (text boxes) : python-docx ne les lit pas via .paragraphs,
    # on les récupère depuis le XML brut (cas fréquent des CV "designés" avec
    # des zones de texte positionnées librement).
    try:
        xml = doc.element.xml
        import re

        textbox_runs = re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml)
        seen: set[str] = set()
        extra: list[str] = []
        existing = set(parts)
        for t in textbox_runs:
            t = t.strip()
            if t and t not in seen and t not in existing:
                seen.add(t)
                extra.append(t)
        parts.extend(extra)
    except Exception:
        pass

    # `doc.inline_shapes` ne détecte pas les images flottantes/ancrées (position
    # libre), fréquentes dans les CV "designés" (ex: un visuel plein-page collé
    # en arrière-plan). On compte directement les fichiers média du .docx
    # (c'est un zip), ce qui est fiable quel que soit le mode d'ancrage.
    n_images = _count_docx_media(path)
    return "\n".join(parts), n_images


def _count_docx_media(path: Path) -> int:
    import zipfile

    with zipfile.ZipFile(path) as z:
        return sum(1 for n in z.namelist() if n.startswith("word/media/"))


def _extract_pptx(path: Path) -> tuple[str, int]:
    prs = Presentation(path)
    parts = []
    n_images = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        parts.append(line)
            if shape.shape_type == 13:  # PICTURE
                n_images += 1
    return "\n".join(parts), n_images


def _render_pdf_pages_as_images(path: Path, out_dir: Path) -> list[Path]:
    doc = fitz.open(path)
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=200)
        out_path = out_dir / f"{path.stem}_p{i}.png"
        pix.save(out_path)
        paths.append(out_path)
    return paths


def _extract_docx_images(path: Path, out_dir: Path, *, min_bytes: int = 3000) -> list[Path]:
    """Extrait les images pertinentes d'un .docx pour l'envoi au modèle vision.

    Deux filtres pour limiter le coût et le bruit envoyés au modèle :
    - déduplication par contenu (une image répétée en en-tête/pied de page
      ne doit être envoyée qu'une fois) ;
    - exclusion des très petites images (`min_bytes`), presque toujours des
      icônes (téléphone, email, LinkedIn) qui n'apportent aucune information
      utile à l'extraction.
    """
    import hashlib
    import zipfile

    out_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    seen_hashes: set[str] = set()

    with zipfile.ZipFile(path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        for name in media:
            data = z.read(name)
            if len(data) < min_bytes:
                continue
            digest = hashlib.sha1(data).hexdigest()
            if digest in seen_hashes:
                continue
            seen_hashes.add(digest)
            out_path = out_dir / f"{path.stem}_{digest[:8]}{Path(name).suffix}"
            out_path.write_bytes(data)
            paths.append(out_path)
    return paths


def extract(path: Path) -> ExtractionResult:
    """Extrait le texte d'un CV. Détecte les cas où le contenu utile est dans
    une image et signale `mode="needs_vision"` avec les images candidates."""
    settings = get_settings()
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        text, n_images = _extract_pdf(path)
    elif suffix == ".docx":
        text, n_images = _extract_docx(path)
    elif suffix == ".pptx":
        text, n_images = _extract_pptx(path)
    else:
        raise ValueError(f"Format non supporté : {suffix}")

    too_short = len(text.strip()) < settings.min_text_chars_for_text_mode
    if too_short and n_images > 0:
        cache_img_dir = settings.cache_dir / "vision_pages"
        if suffix == ".pdf":
            images = _render_pdf_pages_as_images(path, cache_img_dir)
        elif suffix == ".docx":
            images = _extract_docx_images(path, cache_img_dir)
        else:
            images = []  # PPTX : rare en pratique, texte quasi toujours présent
        if images:
            return ExtractionResult(text=text, mode="needs_vision", image_paths=images)

    return ExtractionResult(text=text, mode="text", image_paths=[])
